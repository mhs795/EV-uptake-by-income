"""Write EV_uptake_by_income.xlsx from the processed tables.

Data sheets hold the counts as values; every share, ratio, quintile total and
imputed count is an Excel formula, so changing an input on the Inputs sheet
(e.g. the value given to a suppressed NSW "<=5" cell) flows through every
table and chart.

Run:  python3 build_workbook.py   (after build_data.py)
"""
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

HERE = Path(__file__).resolve().parent
CFG = yaml.safe_load(open(HERE / "config.yaml"))
P = HERE / CFG["paths"]["processed"]
NG = CFG["income"]["n_groups"]
GROUPS = list(range(1, NG + 1))

# ---- styling ---------------------------------------------------------------
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=10, italic=True, color="52514E")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
FILL_HEAD = PatternFill("solid", fgColor="184F95")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_BAND = PatternFill("solid", fgColor="EEF4FC")
THIN = Side(style="thin", color="C8C7C2")
PCT = "0.0%"
NUM = "#,##0"
DEC = "0.00"
MON = "mmm-yy"
USD = "$#,##0"

# ordinal blue ramp for income groups (low -> high), categorical slots for states
GROUP_COLOURS = ["90CAF9", "5AA2EE", "1F7AE0", "1565C0", "0D47A1"]
STATE_COLOURS = {"NSW": "1F7AE0", "QLD": "F57C00", "VIC": "00897B"}
GROUP_LABEL = {g: f"Q{g}" + (" (lowest)" if g == 1 else " (highest)" if g == NG else "") for g in GROUPS}


def head(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42


def put(ws, row, values):
    """Write a row at an explicit row number (append() would ignore our row counter)."""
    for j, v in enumerate(values, 1):
        if v is not None:
            ws.cell(row, j, v)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[L(i)].width = x


def fmt(ws, rng, number_format=None, font=None):
    for row in ws[rng]:
        for c in row:
            if number_format:
                c.number_format = number_format
            if font:
                c.font = font


def base_font(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font == Font():
                    c.font = F_BASE


# ---- data ------------------------------------------------------------------
lga = pd.read_csv(P / "lga_income.csv")
pcode = pd.read_csv(P / "postcode_income.csv")
flow = pd.read_csv(P / "flow_lga_month.csv")
snsw = pd.read_csv(P / "stock_nsw_lga.csv")
sqld = pd.read_csv(P / "stock_qld_proxy_lga.csv")
vic = pd.read_csv(P / "vic_postcode_quarter.csv")
supp = pd.read_csv(P / "suppression.csv")
fuel = pd.read_csv(P / "fuel_prices.csv")
fuel["month"] = pd.to_datetime(fuel["month"])
onset = pd.read_csv(P / "fuel_crisis_onset.csv").iloc[0]
adjustments = pd.read_csv(P / "adjustments.csv").fillna("")

grp = dict(zip(zip(lga.state, lga.lga_name), lga.income_group))
flow["month"] = pd.to_datetime(flow["month"])
flow["group"] = [grp[(s, n)] for s, n in zip(flow.state, flow.lga_name)]
snsw["month"] = pd.to_datetime(snsw["month"])
snsw["group"] = [grp[("NSW", n)] for n in snsw.lga_name]
sqld["month"] = pd.to_datetime(sqld["month"])
sqld["group"] = [grp[("QLD", n)] for n in sqld.lga_name]
vic["qdate"] = [pd.Period(q, "Q").asfreq("M", "end").to_timestamp() for q in vic.quarter]
vgrp = dict(zip(pcode.postcode, pcode.income_group))
vic["group"] = vic.postcode.map(vgrp)

# QLD proxy stock at the same quarter ends as NSW, plus the latest month
q_months = sorted(sqld.month.unique())
q_keep = [m for m in q_months if pd.Timestamp(m).month in CFG["nsw"]["snapshot_months_of_year"]] + [q_months[-1]]
sqld_q = sqld[sqld.month.isin(sorted(set(q_keep)))]

wb = Workbook()

# ============================================================== Inputs ======
wi = wb.active
wi.title = "Inputs"
title(wi, "Inputs and assumptions",
      "Blue text on yellow = assumption you can change. Black = formula. Everything else in the workbook recalculates from these.")
head(wi, 4, ["Item", "Value", "How it was set"])
k = supp.set_index(["table", "fuel_group"])
inputs = [
    ("NSW new regos: value of a suppressed '<=5' cell — BEV", k.loc[("flow", "bev"), "k"], k.loc[("flow", "bev"), "method"]),
    ("NSW new regos: value of a suppressed '<=5' cell — PHEV", k.loc[("flow", "phev"), "k"], k.loc[("flow", "phev"), "method"]),
    ("NSW new regos: value of a suppressed '<=5' cell — all other fuels", k.loc[("flow", "other"), "k"], k.loc[("flow", "other"), "method"]),
    ("NSW fleet: value of a suppressed '<=5' cell — BEV", k.loc[("stock", "bev"), "k"], k.loc[("stock", "bev"), "method"]),
    ("NSW fleet: value of a suppressed '<=5' cell — PHEV", k.loc[("stock", "phev"), "k"], k.loc[("stock", "phev"), "method"]),
    ("NSW fleet: value of a suppressed '<=5' cell — all other fuels", k.loc[("stock", "other"), "k"], k.loc[("stock", "other"), "method"]),
    ("Recent window length (months)", CFG["period"]["recent_window_months"], "config.yaml period.recent_window_months"),
    ("Minimum new private regos in window for an LGA to appear on the scatter", CFG["analysis"]["min_new_regs_scatter"], "config.yaml analysis.min_new_regs_scatter"),
    ("Minimum recent-model vehicles for a VIC postcode to appear on the scatter", CFG["analysis"]["min_recent_vehicles_vic"], "config.yaml analysis.min_recent_vehicles_vic"),
]
for i, (a, b, c) in enumerate(inputs, 5):
    wi.cell(i, 1, a)
    v = wi.cell(i, 2, b)
    v.font, v.fill = F_INPUT, FILL_INPUT
    v.number_format = DEC if isinstance(b, float) else NUM
    wi.cell(i, 3, c)
IN = {"kf_bev": "Inputs!$B$5", "kf_phev": "Inputs!$B$6", "kf_other": "Inputs!$B$7",
      "ks_bev": "Inputs!$B$8", "ks_phev": "Inputs!$B$9", "ks_other": "Inputs!$B$10",
      "window": "Inputs!$B$11", "min_lga": "Inputs!$B$12", "min_vic": "Inputs!$B$13"}

head(wi, 15, ["Derived date", "Value", "Formula"])
derived = [
    ("Latest month of new-registration data — NSW", '=_xlfn.MAXIFS(Flow_LGA_Month!$C:$C,Flow_LGA_Month!$A:$A,"NSW")'),
    ("Latest month of new-registration data — QLD", '=_xlfn.MAXIFS(Flow_LGA_Month!$C:$C,Flow_LGA_Month!$A:$A,"QLD")'),
    ("Recent window start — NSW", f"=EDATE(B16,-({IN['window']}-1))"),
    ("Recent window start — QLD", f"=EDATE(B17,-({IN['window']}-1))"),
    ("First month of new-registration data — both states", "=MIN(Flow_LGA_Month!$C:$C)"),
    ("Baseline window end (first window of the same length)", f"=EDATE(B20,{IN['window']}-1)"),
    ("Latest NSW fleet snapshot", "=MAX(Stock_NSW_LGA!$B:$B)"),
    ("Latest QLD BEV-seen month", "=MAX(Stock_QLD_LGA!$B:$B)"),
    ("Latest VIC fleet snapshot (quarter)", "=MAX(VIC_Postcode_Qtr!$B:$B)"),
    ("VIC snapshot one year earlier", "=EDATE(B24,-12)"),
]
for i, (a, f) in enumerate(derived, 16):
    wi.cell(i, 1, a)
    c = wi.cell(i, 2, f)
    c.number_format = MON
    wi.cell(i, 3, f[1:])
D = {"nsw_last": "Inputs!$B$16", "qld_last": "Inputs!$B$17", "nsw_start": "Inputs!$B$18",
     "qld_start": "Inputs!$B$19", "first": "Inputs!$B$20", "base_end": "Inputs!$B$21",
     "nsw_stock": "Inputs!$B$22", "qld_stock": "Inputs!$B$23", "vic_last": "Inputs!$B$24", "vic_prev": "Inputs!$B$25"}
head(wi, 27, ["Fuel crisis window", "Value", "How it was set"])
c = wi.cell(28, 1, "Crisis start month (2026 Middle East / Strait of Hormuz fuel crisis)")
v = wi.cell(28, 2, pd.Timestamp(onset.onset_month).to_pydatetime())
v.font, v.fill, v.number_format = F_INPUT, FILL_INPUT, MON
wi.cell(28, 3, f"Derived: first month {onset.onset_series} was >= {onset.threshold:.0%} above its trailing 12-month mean "
               f"({onset.onset_price} vs {onset.trailing_mean} c/L). config.yaml fuel.*")
crisis_rows = [
    ("Crisis window end (latest month with data in both states)", f"=MIN({D['nsw_last']},{D['qld_last']})"),
    ("Same months a year earlier — start", "=EDATE(B28,-12)"),
    ("Same months a year earlier — end", "=EDATE(B29,-12)"),
    ("Crisis window length (months)", "=(YEAR(B29)-YEAR(B28))*12+MONTH(B29)-MONTH(B28)+1"),
    ("Pre-crisis window of the same length — start", "=EDATE(B28,-B32)"),
    ("Pre-crisis window — end", "=EDATE(B28,-1)"),
]
for i, (a, f) in enumerate(crisis_rows, 29):
    wi.cell(i, 1, a)
    c = wi.cell(i, 2, f)
    c.number_format = "0" if a.startswith("Crisis window length") else MON
    wi.cell(i, 3, f[1:])
D.update({"cr_start": "Inputs!$B$28", "cr_end": "Inputs!$B$29", "py_start": "Inputs!$B$30", "py_end": "Inputs!$B$31",
          "cr_len": "Inputs!$B$32", "pre_start": "Inputs!$B$33", "pre_end": "Inputs!$B$34"})
widths(wi, [70, 14, 110])

# ======================================================= data sheets ========
wf = wb.create_sheet("Flow_LGA_Month")
cols = ["State", "LGA", "Month", "Income group", "BEV exact", "BEV suppressed cells", "PHEV exact",
        "PHEV suppressed cells", "Other exact", "Other suppressed cells", "BEV (est.)", "PHEV (est.)",
        "Other (est.)", "New private regos (est.)"]
head(wf, 1, cols)
flow = flow.sort_values(["state", "lga_name", "month"])
for r, x in enumerate(flow.itertuples(index=False), 2):
    wf.append([x.state, x.lga_name, x.month.to_pydatetime(), x.group, x.bev_exact, x.bev_supp, x.phev_exact,
               x.phev_supp, x.other_exact, x.other_supp,
               f"=E{r}+{IN['kf_bev']}*F{r}", f"=G{r}+{IN['kf_phev']}*H{r}", f"=I{r}+{IN['kf_other']}*J{r}",
               f"=K{r}+L{r}+M{r}"])
NF = len(flow) + 1
fmt(wf, f"C2:C{NF}", MON)
fmt(wf, f"K2:N{NF}", "#,##0.0")
widths(wf, [7, 26, 10, 9] + [11] * 10)
wf.freeze_panes = "A2"

ws_ = wb.create_sheet("Stock_NSW_LGA")
head(ws_, 1, ["LGA", "Month", "Income group", "BEV exact", "BEV suppressed cells", "PHEV exact",
              "PHEV suppressed cells", "Other exact", "Other suppressed cells", "BEV (est.)", "PHEV (est.)",
              "Other (est.)", "Light vehicles (est.)"])
snsw = snsw.sort_values(["lga_name", "month"])
for r, x in enumerate(snsw.itertuples(index=False), 2):
    ws_.append([x.lga_name, x.month.to_pydatetime(), x.group, x.bev_exact, x.bev_supp, x.phev_exact, x.phev_supp,
                x.other_exact, x.other_supp, f"=D{r}+{IN['ks_bev']}*E{r}", f"=F{r}+{IN['ks_phev']}*G{r}",
                f"=H{r}+{IN['ks_other']}*I{r}", f"=J{r}+K{r}+L{r}"])
NS = len(snsw) + 1
fmt(ws_, f"B2:B{NS}", MON)
fmt(ws_, f"J2:M{NS}", NUM)
widths(ws_, [26, 10, 9] + [11] * 10)
ws_.freeze_panes = "A2"

wq = wb.create_sheet("Stock_QLD_LGA")
head(wq, 1, ["LGA", "Month", "Income group", "BEVs seen since Jan 2022, at last known LGA"])
sqld = sqld.sort_values(["lga_name", "month"])
for x in sqld.itertuples(index=False):
    wq.append([x.lga_name, x.month.to_pydatetime(), x.group, x.bev_seen])
NQ = len(sqld) + 1
fmt(wq, f"B2:B{NQ}", MON)
widths(wq, [26, 10, 9, 22])
wq.freeze_panes = "A2"

wv = wb.create_sheet("VIC_Postcode_Qtr")
head(wv, 1, ["Postcode", "Quarter (last month)", "Income group", "Vehicles", "BEV", "Hybrid (HEV+PHEV)",
             "Recent-model vehicles", "Recent-model BEV"])
vic = vic.sort_values(["postcode", "qdate"])
for x in vic.itertuples(index=False):
    wv.append([x.postcode, x.qdate.to_pydatetime(), x.group, x.vehicles, x.bev, x.hybrid, x.recent_vehicles, x.recent_bev])
NV = len(vic) + 1
fmt(wv, f"B2:B{NV}", MON)
fmt(wv, f"D2:H{NV}", NUM)
widths(wv, [10, 12, 9, 11, 9, 11, 12, 11])
wv.freeze_panes = "A2"

wfp = wb.create_sheet("Fuel_Prices")
FSER = CFG["fuel"]["retail_series"] + CFG["fuel"]["tgp_series"]
FLAB = {"NSW_ULP": "NSW retail ULP", "NSW_Diesel": "NSW retail diesel", "QLD_ULP": "QLD retail ULP",
        "QLD_Diesel": "QLD retail diesel", "TGP_petrol_national": "Terminal gate petrol (national)",
        "TGP_diesel_national": "Terminal gate diesel (national)"}
head(wfp, 1, ["Month"] + [f"{FLAB.get(x, x)} (c/L)" for x in FSER])
for x in fuel.itertuples(index=False):
    wfp.append([x.month.to_pydatetime()] + [round(getattr(x, c), 2) for c in FSER])
NFP = len(fuel) + 1
fmt(wfp, f"A2:A{NFP}", MON)
fmt(wfp, f"B2:{L(1 + len(FSER))}{NFP}", "0.0")
widths(wfp, [10] + [16] * len(FSER))
wfp.freeze_panes = "B2"
FCOL = {x: L(2 + i) for i, x in enumerate(FSER)}

# ========================================================= LGA summary =====
wl = wb.create_sheet("LGA_Summary")
title(wl, "LGA summary — NSW and QLD",
      "Income: ABS Personal Income in Australia 2022-23 (ATO-based), median total income of earners. "
      "New regos = new private vehicles in the recent window. Fleet = latest snapshot.")
lh = ["State", "LGA code", "LGA", "Median income ($)", "Earners", "Income group",
      "New private regos (recent window)", "BEV (recent window)", "PHEV (recent window, NSW)", "BEV share of new",
      "New private regos (baseline window)", "BEV (baseline window)", "BEV share (baseline)", "Change (pp)",
      "BEV fleet (latest)", "Light-vehicle fleet (NSW)", "BEV per 1,000 light vehicles (NSW)", "BEV fleet per 1,000 earners",
      "On scatter? (1 = yes)", "BEV share — fuel crisis months", "BEV share — same months a year earlier",
      "Change during crisis (pp)"]
head(wl, 4, lh)
# eligible LGAs first within each state so the scatter series are contiguous
fl = flow.copy()
fl["new"] = (fl.bev_exact + supp.set_index(["table", "fuel_group"]).loc[("flow", "bev"), "k"] * fl.bev_supp
             + fl.phev_exact + fl.other_exact + fl.phev_supp + fl.other_supp)
last = fl.groupby("state").month.transform("max")
recent = fl[fl.month > last - pd.DateOffset(months=CFG["period"]["recent_window_months"])].groupby(["state", "lga_name"]).new.sum()
lga["elig"] = [recent.get((s, n), 0) >= CFG["analysis"]["min_new_regs_scatter"] for s, n in zip(lga.state, lga.lga_name)]
lga_sorted = lga.sort_values(["state", "elig", "median_income"], ascending=[True, False, True]).reset_index(drop=True)
R0 = 5
for i, x in enumerate(lga_sorted.itertuples(index=False)):
    r = R0 + i
    st = f"IF(A{r}=\"NSW\",{D['nsw_start']},{D['qld_start']})"
    crit = f'Flow_LGA_Month!$A:$A,A{r},Flow_LGA_Month!$B:$B,C{r}'
    rw = f',Flow_LGA_Month!$C:$C,">="&{st}'
    bw = f',Flow_LGA_Month!$C:$C,"<="&{D["base_end"]}'
    stock_bev = (f'IF(A{r}="NSW",SUMIFS(Stock_NSW_LGA!$J:$J,Stock_NSW_LGA!$A:$A,C{r},Stock_NSW_LGA!$B:$B,{D["nsw_stock"]}),'
                 f'SUMIFS(Stock_QLD_LGA!$D:$D,Stock_QLD_LGA!$A:$A,C{r},Stock_QLD_LGA!$B:$B,{D["qld_stock"]}))')
    wl.append([
        x.state, x.lga_code, x.lga_name, x.median_income, x.earners, x.income_group,
        f"=SUMIFS(Flow_LGA_Month!$N:$N,{crit}{rw})",
        f"=SUMIFS(Flow_LGA_Month!$K:$K,{crit}{rw})",
        f'=IF(A{r}="NSW",SUMIFS(Flow_LGA_Month!$L:$L,{crit}{rw}),"n/a")',
        f'=IF(G{r}>0,H{r}/G{r},"")',
        f"=SUMIFS(Flow_LGA_Month!$N:$N,{crit}{bw})",
        f"=SUMIFS(Flow_LGA_Month!$K:$K,{crit}{bw})",
        f'=IF(K{r}>0,L{r}/K{r},"")',
        f'=IF(AND(G{r}>0,K{r}>0),(J{r}-M{r})*100,"")',
        f"={stock_bev}",
        f'=IF(A{r}="NSW",SUMIFS(Stock_NSW_LGA!$M:$M,Stock_NSW_LGA!$A:$A,C{r},Stock_NSW_LGA!$B:$B,{D["nsw_stock"]}),"n/a")',
        f'=IF(A{r}="NSW",IF(P{r}>0,O{r}/P{r}*1000,""),"n/a")',
        f"=O{r}/E{r}*1000",
        f"=IF(G{r}>={IN['min_lga']},1,0)",
        f'=IFERROR(SUMIFS(Flow_LGA_Month!$K:$K,{crit},Flow_LGA_Month!$C:$C,">="&{D["cr_start"]},Flow_LGA_Month!$C:$C,"<="&{D["cr_end"]})/'
        f'SUMIFS(Flow_LGA_Month!$N:$N,{crit},Flow_LGA_Month!$C:$C,">="&{D["cr_start"]},Flow_LGA_Month!$C:$C,"<="&{D["cr_end"]}),"")',
        f'=IFERROR(SUMIFS(Flow_LGA_Month!$K:$K,{crit},Flow_LGA_Month!$C:$C,">="&{D["py_start"]},Flow_LGA_Month!$C:$C,"<="&{D["py_end"]})/'
        f'SUMIFS(Flow_LGA_Month!$N:$N,{crit},Flow_LGA_Month!$C:$C,">="&{D["py_start"]},Flow_LGA_Month!$C:$C,"<="&{D["py_end"]}),"")',
        f'=IF(AND(ISNUMBER(T{r}),ISNUMBER(U{r})),(T{r}-U{r})*100,"")',
    ])
LN = R0 + len(lga_sorted) - 1
fmt(wl, f"D{R0}:D{LN}", USD)
fmt(wl, f"E{R0}:E{LN}", NUM)
fmt(wl, f"G{R0}:I{LN}", NUM)
fmt(wl, f"J{R0}:J{LN}", PCT)
fmt(wl, f"K{R0}:L{LN}", NUM)
fmt(wl, f"M{R0}:M{LN}", PCT)
fmt(wl, f"N{R0}:N{LN}", "0.0")
fmt(wl, f"O{R0}:P{LN}", NUM)
fmt(wl, f"Q{R0}:R{LN}", "0.0")
fmt(wl, f"T{R0}:U{LN}", PCT)
fmt(wl, f"V{R0}:V{LN}", "0.0")
widths(wl, [7, 9, 26, 12, 10, 8, 12, 10, 10, 10, 12, 10, 10, 9, 10, 12, 12, 12, 9, 11, 11, 10])
wl.freeze_panes = "D5"
wl.auto_filter.ref = f"A4:V{LN}"
# scatter ranges (rows computed from the eligibility used to sort)
scat = {}
for s in ("NSW", "QLD"):
    idx = lga_sorted.index[(lga_sorted.state == s) & lga_sorted.elig]
    scat[s] = (R0 + idx.min(), R0 + idx.max())


# ====================================================== VIC postcode ========
wp = wb.create_sheet("VIC_Postcode")
title(wp, "VIC postcodes — fleet snapshot (quarterly) against ATO median taxable income",
      "Income: ATO Taxation Statistics 2023-24, Individuals Table 8. 'Recent-model' = year of manufacture within one year of the snapshot year — a proxy for recent new-vehicle take-up.")
suburbs = dict(pd.read_csv(P / "vic_postcode_suburbs.csv").values)
suburbs.update({int(k): v for k, v in CFG["map"]["postcode_name_overrides"].items()})
ph = ["Postcode", "Median taxable income ($)", "Individuals", "Income group", "Vehicles (latest)", "BEV (latest)",
      "BEV per 1,000 vehicles", "Recent-model vehicles", "Recent-model BEV", "BEV share of recent-model",
      "BEV one year earlier", "Vehicles one year earlier", "BEV added in year per 1,000 vehicles", "On scatter? (1 = yes)",
      "Suburbs (ABS localities in postcode)", "BEV added in latest quarter (crisis)", "Added per 1,000 vehicles (latest quarter)",
      "BEV added same quarter a year earlier", "Added per 1,000 vehicles (year earlier)"]
head(wp, 4, ph)
vl = vic[vic.qdate == vic.qdate.max()].set_index("postcode")
pcode["elig"] = [(vl.recent_vehicles.get(p, 0) >= CFG["analysis"]["min_recent_vehicles_vic"])
                 and (i >= CFG["analysis"]["min_individuals_vic_postcode"]) for p, i in zip(pcode.postcode, pcode.individuals)]
pcode_sorted = pcode.sort_values(["elig", "median_income"], ascending=[False, True]).reset_index(drop=True)
for i, x in enumerate(pcode_sorted.itertuples(index=False)):
    r = R0 + i
    c = f"VIC_Postcode_Qtr!$A:$A,A{r},VIC_Postcode_Qtr!$B:$B,{D['vic_last']}"
    c0 = f"VIC_Postcode_Qtr!$A:$A,A{r},VIC_Postcode_Qtr!$B:$B,{D['vic_prev']}"
    wp.append([
        x.postcode, x.median_income, x.individuals, x.income_group,
        f"=SUMIFS(VIC_Postcode_Qtr!$D:$D,{c})", f"=SUMIFS(VIC_Postcode_Qtr!$E:$E,{c})",
        f'=IF(E{r}>0,F{r}/E{r}*1000,"")',
        f"=SUMIFS(VIC_Postcode_Qtr!$G:$G,{c})", f"=SUMIFS(VIC_Postcode_Qtr!$H:$H,{c})",
        f'=IF(H{r}>0,I{r}/H{r},"")',
        f"=SUMIFS(VIC_Postcode_Qtr!$E:$E,{c0})", f"=SUMIFS(VIC_Postcode_Qtr!$D:$D,{c0})",
        f'=IF(L{r}>0,(F{r}-K{r})/L{r}*1000,"")',
        f"=IF(AND(H{r}>={IN['min_vic']},C{r}>={CFG['analysis']['min_individuals_vic_postcode']}),1,0)",
        suburbs.get(x.postcode, f"Postcode {x.postcode}"),
        f"=F{r}-SUMIFS(VIC_Postcode_Qtr!$E:$E,VIC_Postcode_Qtr!$A:$A,A{r},VIC_Postcode_Qtr!$B:$B,EDATE({D['vic_last']},-3))",
        f'=IF(E{r}>0,P{r}/E{r}*1000,"")',
        f"=SUMIFS(VIC_Postcode_Qtr!$E:$E,VIC_Postcode_Qtr!$A:$A,A{r},VIC_Postcode_Qtr!$B:$B,EDATE({D['vic_last']},-12))"
        f"-SUMIFS(VIC_Postcode_Qtr!$E:$E,VIC_Postcode_Qtr!$A:$A,A{r},VIC_Postcode_Qtr!$B:$B,EDATE({D['vic_last']},-15))",
        f'=IF(L{r}>0,R{r}/L{r}*1000,"")',
    ])
PN = R0 + len(pcode_sorted) - 1
fmt(wp, f"B{R0}:B{PN}", USD)
fmt(wp, f"C{R0}:C{PN}", NUM)
fmt(wp, f"E{R0}:F{PN}", NUM)
fmt(wp, f"G{R0}:G{PN}", "0.0")
fmt(wp, f"H{R0}:I{PN}", NUM)
fmt(wp, f"J{R0}:J{PN}", PCT)
fmt(wp, f"K{R0}:L{PN}", NUM)
fmt(wp, f"M{R0}:M{PN}", "0.0")
fmt(wp, f"P{R0}:P{PN}", NUM)
fmt(wp, f"Q{R0}:Q{PN}", "0.0")
fmt(wp, f"R{R0}:R{PN}", NUM)
fmt(wp, f"S{R0}:S{PN}", "0.0")

widths(wp, [10, 13, 11, 8, 11, 9, 11, 11, 10, 11, 11, 12, 13, 9])
wp.freeze_panes = "B5"
wp.auto_filter.ref = f"A4:S{PN}"
wp.column_dimensions["O"].width = 34
vic_elig = pcode_sorted.index[pcode_sorted.elig]
scat["VIC"] = (R0 + vic_elig.min(), R0 + vic_elig.max())


# ============================================ Flow by income group, monthly ==
wm = wb.create_sheet("Flow_Group_Month")
title(wm, "New private registrations by LGA income group — monthly",
      "Groups are earner-weighted: each holds about a fifth of the state's earners. Q1 = lowest-income LGAs. Counts are SUMIFS over Flow_LGA_Month.")
months = sorted(flow.month.unique())
blocks = [("NSW", "new", "N"), ("NSW", "bev", "K"), ("QLD", "new", "N"), ("QLD", "bev", "K")]
hdr = ["Month"]
for s, m, _ in blocks:
    hdr += [f"{s} {'new' if m == 'new' else 'BEV'} {GROUP_LABEL[g]}" for g in GROUPS]
share_start = 1 + len(blocks) * NG + 1
for s in ("NSW", "QLD"):
    hdr += [f"{s} BEV share {GROUP_LABEL[g]}" for g in GROUPS] + [f"{s} BEV share all"]
head(wm, 4, hdr)
M0 = 5
for i, m in enumerate(months):
    r = M0 + i
    wm.cell(r, 1, pd.Timestamp(m).to_pydatetime()).number_format = MON
    col = 2
    for s, _, src in blocks:
        for g in GROUPS:
            wm.cell(r, col, f'=SUMIFS(Flow_LGA_Month!${src}:${src},Flow_LGA_Month!$A:$A,"{s}",'
                            f'Flow_LGA_Month!$D:$D,{g},Flow_LGA_Month!$C:$C,$A{r})').number_format = NUM
            col += 1
    for bi, s in enumerate(("NSW", "QLD")):
        new0 = 2 + bi * 2 * NG
        bev0 = new0 + NG
        for g in GROUPS:
            n_, b_ = L(new0 + g - 1), L(bev0 + g - 1)
            wm.cell(r, col, f'=IF({n_}{r}>0,{b_}{r}/{n_}{r},"")').number_format = PCT
            col += 1
        nr = f"{L(new0)}{r}:{L(new0 + NG - 1)}{r}"
        br = f"{L(bev0)}{r}:{L(bev0 + NG - 1)}{r}"
        wm.cell(r, col, f'=IF(SUM({nr})>0,SUM({br})/SUM({nr}),"")').number_format = PCT
        col += 1
MN = M0 + len(months) - 1
widths(wm, [9] + [9] * (len(hdr) - 1))
wm.freeze_panes = "B5"
SH = {"NSW": share_start, "QLD": share_start + NG + 1}

# ============================================== Flow by income group, window ==
wg = wb.create_sheet("Flow_Group_Summary")
title(wg, "Who is buying the new BEVs? — recent window vs baseline window",
      "Recent window = latest N months (Inputs). Baseline = the first N months of data. "
      "Representation index = group's share of BEVs ÷ its share of all new private regos (1.0 = proportional).")
gh = ["State", "Income group", "Lowest LGA median ($)", "Highest LGA median ($)", "Earners",
      "New private regos", "BEV", "BEV share of new", "Share of state's new BEVs", "Share of state's new regos",
      "Representation index", "BEV per 1,000 earners", "Baseline: new regos", "Baseline: BEV", "Baseline: BEV share",
      "Baseline: representation index"]
head(wg, 4, gh)
G0 = 5
grow = {}
r = G0
for s in ("NSW", "QLD"):
    start = D["nsw_start"] if s == "NSW" else D["qld_start"]
    first = r
    for g in GROUPS + ["All"]:
        is_all = g == "All"
        gc = "" if is_all else f",Flow_LGA_Month!$D:$D,{g}"
        lc = "" if is_all else f",LGA_Summary!$F:$F,{g}"
        rw = f',Flow_LGA_Month!$C:$C,">="&{start}'
        bw = f',Flow_LGA_Month!$C:$C,"<="&{D["base_end"]}'
        allr = first + NG
        put(wg, r, [
            s, "All" if is_all else GROUP_LABEL[g],
            f'=_xlfn.MINIFS(LGA_Summary!$D:$D,LGA_Summary!$A:$A,"{s}"{lc})',
            f'=_xlfn.MAXIFS(LGA_Summary!$D:$D,LGA_Summary!$A:$A,"{s}"{lc})',
            f'=SUMIFS(LGA_Summary!$E:$E,LGA_Summary!$A:$A,"{s}"{lc})',
            f'=SUMIFS(Flow_LGA_Month!$N:$N,Flow_LGA_Month!$A:$A,"{s}"{gc}{rw})',
            f'=SUMIFS(Flow_LGA_Month!$K:$K,Flow_LGA_Month!$A:$A,"{s}"{gc}{rw})',
            f"=G{r}/F{r}", f"=G{r}/G${allr}", f"=F{r}/F${allr}", f"=I{r}/J{r}", f"=G{r}/E{r}*1000",
            f'=SUMIFS(Flow_LGA_Month!$N:$N,Flow_LGA_Month!$A:$A,"{s}"{gc}{bw})',
            f'=SUMIFS(Flow_LGA_Month!$K:$K,Flow_LGA_Month!$A:$A,"{s}"{gc}{bw})',
            f"=N{r}/M{r}", f"=(N{r}/N${allr})/(M{r}/M${allr})",
        ])
        if is_all:
            for c in range(1, 17):
                wg.cell(r, c).font = F_BOLD
        r += 1
    grow[s] = first
    put(wg, r, [s, f"Top ÷ bottom group", None, None, None, None, None, f"=H{first + NG - 1}/H{first}", None, None,
               f"=K{first + NG - 1}/K{first}", f"=L{first + NG - 1}/L{first}", None, None,
               f"=O{first + NG - 1}/O{first}", f"=P{first + NG - 1}/P{first}"])
    for c in range(1, 17):
        wg.cell(r, c).font = F_BOLD
    r += 2
GN = r - 1
fmt(wg, f"C{G0}:D{GN}", USD)
fmt(wg, f"E{G0}:G{GN}", NUM)
fmt(wg, f"H{G0}:J{GN}", PCT)
fmt(wg, f"K{G0}:K{GN}", DEC)
fmt(wg, f"L{G0}:L{GN}", "0.0")
fmt(wg, f"M{G0}:N{GN}", NUM)
fmt(wg, f"O{G0}:O{GN}", PCT)
fmt(wg, f"P{G0}:P{GN}", DEC)
for s in ("NSW", "QLD"):
    for c in ("H", "K", "L", "O", "P"):
        wg[f"{c}{grow[s] + NG + 1}"].number_format = DEC

# VIC block
r += 1
wg.cell(r, 1, "VIC — postcode income groups (ATO 2023-24). Flow proxy = BEV share of recent-model vehicles in the latest snapshot; "
              "fleet additions = BEV fleet growth over the last year.").font = F_BOLD
r += 1
vh = ["State", "Income group", "Lowest postcode median ($)", "Highest postcode median ($)", "Individuals",
      "Recent-model vehicles", "Recent-model BEV", "BEV share of recent-model", "Share of VIC recent-model BEVs",
      "Share of VIC recent-model vehicles", "Representation index", "Recent-model BEV per 1,000 individuals",
      "BEV added in last year", "BEV added per 1,000 individuals"]
head(wg, r, vh)
r += 1
vfirst = r
for g in GROUPS + ["All"]:
    is_all = g == "All"
    gc = "" if is_all else f",VIC_Postcode!$D:$D,{g}"
    allr = vfirst + NG
    put(wg, r, [
        "VIC", "All" if is_all else GROUP_LABEL[g],
        f'=_xlfn.MINIFS(VIC_Postcode!$B:$B,VIC_Postcode!$A:$A,">0"{gc})',
        f'=_xlfn.MAXIFS(VIC_Postcode!$B:$B,VIC_Postcode!$A:$A,">0"{gc})',
        f'=SUMIFS(VIC_Postcode!$C:$C,VIC_Postcode!$A:$A,">0"{gc})',
        f'=SUMIFS(VIC_Postcode!$H:$H,VIC_Postcode!$A:$A,">0"{gc})',
        f'=SUMIFS(VIC_Postcode!$I:$I,VIC_Postcode!$A:$A,">0"{gc})',
        f"=G{r}/F{r}", f"=G{r}/G${allr}", f"=F{r}/F${allr}", f"=I{r}/J{r}", f"=G{r}/E{r}*1000",
        f'=SUMIFS(VIC_Postcode!$F:$F,VIC_Postcode!$A:$A,">0"{gc})-SUMIFS(VIC_Postcode!$K:$K,VIC_Postcode!$A:$A,">0"{gc})',
        f"=M{r}/E{r}*1000",
    ])
    if is_all:
        for c in range(1, 15):
            wg.cell(r, c).font = F_BOLD
    r += 1
grow["VIC"] = vfirst
put(wg, r, ["VIC", "Top ÷ bottom group", None, None, None, None, None, f"=H{vfirst + NG - 1}/H{vfirst}", None, None,
           f"=K{vfirst + NG - 1}/K{vfirst}", f"=L{vfirst + NG - 1}/L{vfirst}", None, f"=N{vfirst + NG - 1}/N{vfirst}"])
for c in range(1, 15):
    wg.cell(r, c).font = F_BOLD
VN = r
fmt(wg, f"C{vfirst}:D{VN}", USD)
fmt(wg, f"E{vfirst}:G{VN}", NUM)
fmt(wg, f"H{vfirst}:J{VN}", PCT)
fmt(wg, f"K{vfirst}:K{VN}", DEC)
fmt(wg, f"L{vfirst}:L{VN}", "0.0")
fmt(wg, f"M{vfirst}:M{VN}", NUM)
fmt(wg, f"N{vfirst}:N{VN}", "0.0")
for c in ("H", "K", "L", "N"):
    wg[f"{c}{VN}"].number_format = DEC
widths(wg, [7, 14, 12, 12, 11, 12, 10, 10, 11, 11, 11, 11, 11, 10, 10, 11])
wg.freeze_panes = "C5"

# =============================================== Fleet (stock) by group =====
wk = wb.create_sheet("Stock_Group")
title(wk, "Registered fleet (stock) by income group",
      "NSW: TfNSW registration snapshot, light vehicles, quarter ends. QLD: BEVs seen in any transaction since Jan 2022 at their last "
      "known LGA (QLD publishes no regional fleet-by-fuel snapshot). VIC: DTP whole-fleet snapshot by postcode.")
blocks_s = [
    ("NSW", "Light vehicles", "Stock_NSW_LGA", "M", "A", "B", "C", sorted(snsw.month.unique())),
    ("NSW", "BEV", "Stock_NSW_LGA", "J", "A", "B", "C", None),
    ("QLD", "BEV seen", "Stock_QLD_LGA", "D", "A", "B", "C", sorted(sqld_q.month.unique())),
    ("VIC", "Vehicles", "VIC_Postcode_Qtr", "D", "A", "B", "C", sorted(vic.qdate.unique())),
    ("VIC", "BEV", "VIC_Postcode_Qtr", "E", "A", "B", "C", None),
    ("VIC", "Recent-model vehicles", "VIC_Postcode_Qtr", "G", "A", "B", "C", None),
    ("VIC", "Recent-model BEV", "VIC_Postcode_Qtr", "H", "A", "B", "C", None),
]
# layout: three tables stacked, each with its own date column
STK = {}
r = 4


def stock_table(r, state, parts, dates, ratio_specs):
    """parts: list of (label, sheet, value col, group col, date col); ratio_specs: (label, num part, den part, scale)"""
    hdr = ["Month"]
    for lab, *_ in parts:
        hdr += [f"{state} {lab} {GROUP_LABEL[g]}" for g in GROUPS]
    for lab, *_ in ratio_specs:
        hdr += [f"{state} {lab} {GROUP_LABEL[g]}" for g in GROUPS] + [f"{state} {lab} all"]
    head(wk, r, hdr)
    first = r + 1
    for i, d in enumerate(dates):
        rr = first + i
        wk.cell(rr, 1, pd.Timestamp(d).to_pydatetime()).number_format = MON
        col = 2
        for lab, sheet, vc, gc, dc in parts:
            for g in GROUPS:
                wk.cell(rr, col, f"=SUMIFS({sheet}!${vc}:${vc},{sheet}!${gc}:${gc},{g},{sheet}!${dc}:${dc},$A{rr})").number_format = NUM
                col += 1
        for lab, num_i, den, scale in ratio_specs:
            n0 = 2 + num_i * NG
            for g in GROUPS:
                nref = f"{L(n0 + g - 1)}{rr}"
                if isinstance(den, int):
                    dref = f"{L(2 + den * NG + g - 1)}{rr}"
                else:  # earners from LGA_Summary
                    dref = f'SUMIFS(LGA_Summary!$E:$E,LGA_Summary!$A:$A,"{state}",LGA_Summary!$F:$F,{g})'
                wk.cell(rr, col, f"=IF({dref}>0,{nref}/{dref}*{scale},\"\")").number_format = "0.0" if scale > 1 else PCT
                col += 1
            nall = f"SUM({L(n0)}{rr}:{L(n0 + NG - 1)}{rr})"
            if isinstance(den, int):
                dall = f"SUM({L(2 + den * NG)}{rr}:{L(2 + den * NG + NG - 1)}{rr})"
            else:
                dall = f'SUMIFS(LGA_Summary!$E:$E,LGA_Summary!$A:$A,"{state}")'
            wk.cell(rr, col, f"={nall}/{dall}*{scale}").number_format = "0.0" if scale > 1 else PCT
            col += 1
    return first, first + len(dates) - 1, len(parts)


nsw_dates = sorted(snsw.month.unique())
STK["NSW"] = stock_table(r, "NSW", [("light vehicles", "Stock_NSW_LGA", "M", "C", "B"), ("BEV", "Stock_NSW_LGA", "J", "C", "B")],
                         nsw_dates, [("BEV per 1,000 light vehicles", 1, 0, 1000), ("BEV per 1,000 earners", 1, "earners", 1000)])
r = STK["NSW"][1] + 3
STK["QLD"] = stock_table(r, "QLD", [("BEV seen", "Stock_QLD_LGA", "D", "C", "B")],
                         sorted(sqld_q.month.unique()), [("BEV per 1,000 earners", 0, "earners", 1000)])
r = STK["QLD"][1] + 3
STK["VIC"] = stock_table(r, "VIC", [("vehicles", "VIC_Postcode_Qtr", "D", "C", "B"), ("BEV", "VIC_Postcode_Qtr", "E", "C", "B"),
                                    ("recent-model vehicles", "VIC_Postcode_Qtr", "G", "C", "B"),
                                    ("recent-model BEV", "VIC_Postcode_Qtr", "H", "C", "B")],
                         sorted(vic.qdate.unique()), [("BEV per 1,000 vehicles", 1, 0, 1000),
                                                      ("BEV share of recent-model", 3, 2, 1)])
widths(wk, [9] + [10] * 40)
wk.freeze_panes = "B4"

def style_axes(ch, y_fmt=None, x_fmt=None, y_title=None, x_title=None):
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    if y_fmt:
        ch.y_axis.number_format = y_fmt
        ch.y_axis.numFmt.sourceLinked = False if ch.y_axis.numFmt is not None else None
    if x_fmt:
        ch.x_axis.number_format = x_fmt
    if y_title:
        ch.y_axis.title = y_title
    if x_title:
        ch.x_axis.title = x_title
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E4E3DF"))
    if ch.legend is not None:
        ch.legend.position = "b"
    ch.height, ch.width = 10, 17
    small = CharacterProperties(sz=800, solidFill="52514E")
    for ax in (ch.x_axis, ch.y_axis):
        ax.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=small), endParaRPr=small)])
    if ch.title is not None:
        cp = CharacterProperties(sz=1100, b=True, solidFill="0B0B0B")
        for para in ch.title.tx.rich.p:
            para.pPr = ParagraphProperties(defRPr=cp)
            for run in para.r or []:
                run.rPr = cp


def line_series(s, colour, width=22000):
    s.graphicalProperties.line.solidFill = colour
    s.graphicalProperties.line.width = width
    s.marker.symbol = "none"
    s.smooth = False



# ============================================================ Fuel crisis ==
wx = wb.create_sheet("Fuel_Crisis")
title(wx, "The 2026 fuel crisis — did the price shock change who buys BEVs?",
      "Crisis = months from the detected onset (Inputs) to the latest month, compared with the same calendar months a year earlier "
      "(controls for seasonality such as the June EOFY peak).")
# ---- monthly prices next to BEV share
head(wx, 4, ["Month", "NSW retail ULP (c/L)", "NSW retail diesel (c/L)", "QLD retail ULP (c/L)", "QLD retail diesel (c/L)",
             "Terminal gate petrol (c/L)", "Crisis month? (1 = yes)", "NSW BEV share — all areas", "QLD BEV share — all areas"])
X0 = 5
price_cols = ["NSW_ULP", "NSW_Diesel", "QLD_ULP", "QLD_Diesel", "TGP_petrol_national"]
for i in range(len(months)):
    r, fr = X0 + i, M0 + i
    wx.cell(r, 1, f"=Flow_Group_Month!A{fr}").number_format = MON
    for j, sname in enumerate(price_cols, 2):
        wx.cell(r, j, f"=IFERROR(INDEX(Fuel_Prices!${FCOL[sname]}:${FCOL[sname]},MATCH($A{r},Fuel_Prices!$A:$A,0)),\"\")").number_format = "0.0"
    wx.cell(r, 7, f"=IF(AND(A{r}>={D['cr_start']},A{r}<={D['cr_end']}),1,0)")
    wx.cell(r, 8, f"=Flow_Group_Month!{L(SH['NSW'] + NG)}{fr}").number_format = PCT
    wx.cell(r, 9, f"=Flow_Group_Month!{L(SH['QLD'] + NG)}{fr}").number_format = PCT
    for j in (1, 8, 9):
        wx.cell(r, j).font = F_LINK
XN = X0 + len(months) - 1

# ---- by income group
r = XN + 3
wx.cell(r, 1, "New private registrations by income group — crisis months vs same months a year earlier vs pre-crisis months").font = F_BOLD
r += 1
xh = ["State", "Income group", "New regos (crisis)", "BEV (crisis)", "BEV share (crisis)", "New regos (year earlier)",
      "BEV (year earlier)", "BEV share (year earlier)", "Change vs year earlier (pp)", "BEV share multiple (× year earlier)",
      "Extra BEVs vs year earlier", "Share of the extra BEVs", "Share of crisis new regos", "Non-BEV new regos, change vs year earlier"]
head(wx, r, xh)
r += 1
xgrow = {}
for s_ in ("NSW", "QLD"):
    first = r
    allr = first + NG
    for g in GROUPS + ["All"]:
        is_all = g == "All"
        gc = "" if is_all else f",Flow_LGA_Month!$D:$D,{g}"
        base = f'Flow_LGA_Month!$A:$A,"{s_}"{gc}'

        def w(col, a, b):
            return f'SUMIFS(Flow_LGA_Month!${col}:${col},{base},Flow_LGA_Month!$C:$C,">="&{a},Flow_LGA_Month!$C:$C,"<="&{b})'
        put(wx, r, [
            s_, "All" if is_all else GROUP_LABEL[g],
            "=" + w("N", D["cr_start"], D["cr_end"]), "=" + w("K", D["cr_start"], D["cr_end"]), f"=D{r}/C{r}",
            "=" + w("N", D["py_start"], D["py_end"]), "=" + w("K", D["py_start"], D["py_end"]), f"=G{r}/F{r}",
            f"=(E{r}-H{r})*100", f"=E{r}/H{r}",
            f"=D{r}-G{r}", f"=K{r}/K${allr}", f"=C{r}/C${allr}", f"=(C{r}-D{r})/(F{r}-G{r})-1",
        ])
        if is_all:
            for c_ in range(1, 15):
                wx.cell(r, c_).font = F_BOLD
        r += 1
    xgrow[s_] = first
    put(wx, r, [s_, "Top ÷ bottom group", None, None, f"=E{first + NG - 1}/E{first}", None, None, f"=H{first + NG - 1}/H{first}",
               f"=I{first + NG - 1}-I{first}", f"=J{first + NG - 1}/J{first}"])
    for c_ in range(1, 15):
        wx.cell(r, c_).font = F_BOLD
    r += 2
XGN = r
fmt(wx, f"C{xgrow['NSW']}:D{XGN}", NUM)
fmt(wx, f"E{xgrow['NSW']}:E{XGN}", PCT)
fmt(wx, f"F{xgrow['NSW']}:G{XGN}", NUM)
fmt(wx, f"H{xgrow['NSW']}:H{XGN}", PCT)
fmt(wx, f"I{xgrow['NSW']}:I{XGN}", "0.0")
fmt(wx, f"J{xgrow['NSW']}:J{XGN}", DEC)
fmt(wx, f"K{xgrow['NSW']}:K{XGN}", NUM)
fmt(wx, f"L{xgrow['NSW']}:N{XGN}", PCT)
for s_ in ("NSW", "QLD"):
    for c_ in "EHJ":
        wx[f"{c_}{xgrow[s_] + NG + 1}"].number_format = DEC
    wx[f"I{xgrow[s_] + NG + 1}"].number_format = "0.0"

# ---- VIC: fleet additions in the crisis quarter vs a year earlier
r += 1
wx.cell(r, 1, "VIC — BEVs added to the fleet in the latest quarter vs the same quarter a year earlier "
              "(VIC has no monthly regional series; the latest quarter falls inside the crisis)").font = F_BOLD
r += 1
head(wx, r, ["State", "Income group", "BEV fleet latest quarter", "BEV fleet quarter before", "Added (latest quarter)",
             "BEV fleet same quarter a year earlier", "BEV fleet quarter before that", "Added (year earlier)",
             "Multiple (× year earlier)", "Vehicles (latest)", "Added per 1,000 vehicles (latest quarter)",
             "Added per 1,000 vehicles (year earlier)", "Share of added BEVs (latest quarter)"])
r += 1
vx0 = r
qdates = {"q0": D["vic_last"], "q1": f"EDATE({D['vic_last']},-3)", "q4": f"EDATE({D['vic_last']},-12)", "q5": f"EDATE({D['vic_last']},-15)"}
for g in GROUPS + ["All"]:
    is_all = g == "All"
    gc = "" if is_all else f",VIC_Postcode_Qtr!$C:$C,{g}"

    def vq(col, d):
        return f'SUMIFS(VIC_Postcode_Qtr!${col}:${col},VIC_Postcode_Qtr!$B:$B,{d}{gc})'
    allr = vx0 + NG
    put(wx, r, ["VIC", "All" if is_all else GROUP_LABEL[g], "=" + vq("E", qdates["q0"]), "=" + vq("E", qdates["q1"]), f"=C{r}-D{r}",
               "=" + vq("E", qdates["q4"]), "=" + vq("E", qdates["q5"]), f"=F{r}-G{r}", f"=E{r}/H{r}",
               "=" + vq("D", qdates["q0"]), f"=E{r}/J{r}*1000", f"=H{r}/({vq('D', qdates['q4'])})*1000", f"=E{r}/E${allr}"])
    if is_all:
        for c_ in range(1, 14):
            wx.cell(r, c_).font = F_BOLD
    r += 1
put(wx, r, ["VIC", "Top ÷ bottom group", None, None, None, None, None, None, f"=I{vx0 + NG - 1}/I{vx0}", None,
           f"=K{vx0 + NG - 1}/K{vx0}", f"=L{vx0 + NG - 1}/L{vx0}"])
for c_ in range(1, 14):
    wx.cell(r, c_).font = F_BOLD
fmt(wx, f"C{vx0}:H{r}", NUM)
fmt(wx, f"I{vx0}:I{r}", DEC)
fmt(wx, f"J{vx0}:J{r}", NUM)
fmt(wx, f"K{vx0}:L{r}", "0.0")
fmt(wx, f"M{vx0}:M{r}", PCT)
for c_ in "KL":
    wx[f"{c_}{r}"].number_format = DEC
xgrow["VIC"] = vx0
widths(wx, [9, 14] + [12] * 17)
wx.freeze_panes = "C5"

# ============================================================ Raw numbers ===
nf, nl, _ = STK["NSW"]
qf, ql, _ = STK["QLD"]
vf, vl_, _ = STK["VIC"]
wrn = wb.create_sheet("Raw_Numbers")
title(wrn, "Raw numbers — vehicle counts, not shares",
      "New private registrations per month (NSW, QLD) and BEV counts by income group. NSW counts include the estimate for suppressed '<=5' cells (Inputs).")
head(wrn, 4, ["Month", "NSW new private regos", "NSW BEV", "NSW other fuels", "QLD new private regos", "QLD BEV", "QLD other fuels"])
RN0 = 5
for i in range(len(months)):
    r, fr = RN0 + i, M0 + i
    wrn.cell(r, 1, f"=Flow_Group_Month!A{fr}").number_format = MON
    for bi in range(2):
        new0 = 2 + bi * 2 * NG
        bev0 = new0 + NG
        c0 = 2 + bi * 3
        wrn.cell(r, c0, f"=SUM(Flow_Group_Month!{L(new0)}{fr}:{L(new0 + NG - 1)}{fr})").number_format = NUM
        wrn.cell(r, c0 + 1, f"=SUM(Flow_Group_Month!{L(bev0)}{fr}:{L(bev0 + NG - 1)}{fr})").number_format = NUM
        wrn.cell(r, c0 + 2, f"={L(c0)}{r}-{L(c0 + 1)}{r}").number_format = NUM
RNN = RN0 + len(months) - 1

GC = 10  # group table starts at column J
head(wrn, 4, ["Income group", "NSW new BEVs (recent window)", "QLD new BEVs (recent window)", "VIC recent-model BEVs in fleet",
              "NSW BEV fleet (latest)", "QLD BEVs seen since 2022 (latest)", "VIC BEV fleet (latest)",
              "NSW new BEVs — year before crisis", "NSW new BEVs — crisis months",
              "QLD new BEVs — year before crisis", "QLD new BEVs — crisis months"], col=GC)
for gi, g in enumerate(GROUPS):
    r = RN0 + gi
    vals = [GROUP_LABEL[g],
            f"=Flow_Group_Summary!G{grow['NSW'] + gi}", f"=Flow_Group_Summary!G{grow['QLD'] + gi}", f"=Flow_Group_Summary!G{grow['VIC'] + gi}",
            f"=Stock_Group!{L(2 + NG + gi)}{nl}", f"=Stock_Group!{L(2 + gi)}{ql}", f"=Stock_Group!{L(2 + NG + gi)}{vl_}",
            f"=Fuel_Crisis!G{xgrow['NSW'] + gi}", f"=Fuel_Crisis!D{xgrow['NSW'] + gi}",
            f"=Fuel_Crisis!G{xgrow['QLD'] + gi}", f"=Fuel_Crisis!D{xgrow['QLD'] + gi}"]
    for j, v in enumerate(vals):
        c = wrn.cell(r, GC + j, v)
        if j:
            c.number_format, c.font = NUM, F_LINK
RG_ALL = RN0 + NG
wrn.cell(RG_ALL, GC, "All").font = F_BOLD
for j in range(1, 11):
    c = wrn.cell(RG_ALL, GC + j, f"=SUM({L(GC + j)}{RN0}:{L(GC + j)}{RN0 + NG - 1})")
    c.number_format, c.font = NUM, F_BOLD
widths(wrn, [9, 12, 10, 12, 12, 10, 12, 3, 3, 14] + [14] * 10)
wrn.freeze_panes = "B5"

# ============================================================= Top 10s =====
TOPN = CFG["analysis"]["top_n"]
wr = wb.create_sheet("Rank_Keys")
wr["A1"] = "Helper keys for the Top10 sheet: metric value plus a tiny row-number tie-break, or ±1E9 when the area is not eligible. Do not edit."
wt = wb.create_sheet("Top10")
title(wt, f"Top and bottom {TOPN} areas",
      "Suburb-level EV data is not published by any state: VIC is by postcode (named by its ABS suburbs), NSW and QLD by LGA. "
      "Only areas above the size thresholds on Inputs are ranked. Live formulas.")
lga_show = [("LGA", "C", None), ("Median income ($)", "D", USD), ("Income group", "F", "0")]
vic_show = [("Postcode", "A", "0"), ("Suburbs", "O", None), ("Median taxable income ($)", "B", USD), ("Income group", "D", "0")]
specs = [
    ("NSW LGAs — highest BEV share of new private cars (recent window)", "LGA", "NSW", "J", -1, [("New regos", "G", NUM), ("BEV share", "J", PCT)]),
    ("NSW LGAs — lowest BEV share of new private cars (recent window)", "LGA", "NSW", "J", 1, [("New regos", "G", NUM), ("BEV share", "J", PCT)]),
    ("NSW LGAs — most new BEVs registered (recent window, count)", "LGA", "NSW", "H", -1, [("New regos", "G", NUM), ("New BEVs", "H", NUM)]),
    ("NSW LGAs — largest BEV fleet (count)", "LGA", "NSW", "O", -1, [("Per 1,000 vehicles", "Q", "0.0"), ("BEV fleet", "O", NUM)]),
    ("NSW LGAs — biggest rise in BEV share during the fuel crisis", "LGA", "NSW", "V", -1,
     [("Year earlier", "U", PCT), ("Crisis", "T", PCT), ("Change (pp)", "V", "0.0")]),
    ("QLD LGAs — highest BEV share of new private cars (recent window)", "LGA", "QLD", "J", -1, [("New regos", "G", NUM), ("BEV share", "J", PCT)]),
    ("QLD LGAs — most new BEVs registered (recent window, count)", "LGA", "QLD", "H", -1, [("New regos", "G", NUM), ("New BEVs", "H", NUM)]),
    ("VIC suburbs (postcodes) — most BEVs per 1,000 vehicles", "VIC", None, "G", -1, [("BEVs", "F", NUM), ("Per 1,000", "G", "0.0")]),
    ("VIC suburbs (postcodes) — most BEVs registered (fleet count)", "VIC", None, "F", -1, [("Vehicles", "E", NUM), ("BEVs", "F", NUM)]),
    ("VIC suburbs (postcodes) — most BEVs added in the crisis quarter (count)", "VIC", None, "P", -1,
     [("Added year earlier", "R", NUM), ("Added (crisis qtr)", "P", NUM)]),
]
top_tables = []
row = 4
key_col = 1
for k, (ttl, src, st, mcol, direction, extra) in enumerate(specs):
    key_col += 1
    kc = L(key_col)
    sheet, first, last = ("LGA_Summary", R0, LN) if src == "LGA" else ("VIC_Postcode", R0, PN)
    elig = f"{sheet}!$S{{r}}=1" if src == "LGA" else f"{sheet}!$N{{r}}=1"
    stc = f'{sheet}!$A{{r}}="{st}",' if st else ""
    wr.cell(3, key_col, ttl)
    for rr in range(first, last + 1):
        v = f"{sheet}!${mcol}{rr}"
        cond = f"AND({stc.format(r=rr)}{elig.format(r=rr)},ISNUMBER({v}))"
        if direction < 0:
            wr.cell(rr, key_col, f"=IF({cond},{v}+ROW()/1E9,-1E9)")
        else:
            wr.cell(rr, key_col, f"=IF({cond},{v}-ROW()/1E9,1E9)")
    show = (lga_show if src == "LGA" else vic_show) + extra
    col0 = 1 if k % 2 == 0 else 1 + 9
    wt.cell(row, col0, ttl).font = F_BOLD
    head(wt, row + 1, ["#"] + [h for h, _, _ in show], col=col0)
    pick = "LARGE" if direction < 0 else "SMALL"
    for n in range(1, TOPN + 1):
        rr = row + 1 + n
        wt.cell(rr, col0, n)
        mref = f"MATCH({pick}(Rank_Keys!${kc}:${kc},{n}),Rank_Keys!${kc}:${kc},0)"
        for j, (h, c, f) in enumerate(show, 1):
            cell = wt.cell(rr, col0 + j, f"=IFERROR(INDEX({sheet}!${c}:${c},{mref}),\"\")")
            if f:
                cell.number_format = f
            cell.font = F_LINK
    top_tables.append((ttl, row + 2, row + 1 + TOPN, col0 + 1, col0 + len(show)))
    if k % 2 == 1:
        row += TOPN + 4
for c_ in range(1, 20):
    wt.column_dimensions[L(c_)].width = 12
for c_ in ("B", "K"):
    wt.column_dimensions[c_].width = 24
for c_ in ("C", "L"):
    wt.column_dimensions[c_].width = 14
wr.sheet_state = "hidden"

# ============================================================= Suppression ==
wsu = wb.create_sheet("Suppression")
title(wsu, "NSW small-cell suppression — how the '<=5' cells are valued",
      "TfNSW publishes every count of 5 or fewer as '<=5'. Because each row is also split by colour, gender, age group etc., "
      "almost every row is suppressed, so the value given to a '<=5' cell sets the level of every NSW count.")
head(wsu, 4, ["Table", "Fuel group", "Estimated value per '<=5' cell", "Method"])
for x in supp.itertuples(index=False):
    wsu.append([x.table, x.fuel_group, x.k, x.method])
fmt(wsu, "C5:C10", DEC)
notes = [
    "Why it matters less than it looks: every NSW comparison in this workbook is a ratio within NSW (BEV share of new regos, BEVs per 1,000 vehicles), "
    "so a common scaling of all cells cancels. Only a difference between the size of BEV cells and other cells moves the income gradient.",
    "Flow estimate: QLD publishes unit records. Each QLD new private vehicle was given a synthetic gender x age group drawn from the NSW private "
    "new-vehicle mix for its fuel group, then aggregated to NSW's grain (month x LGA x make x fuel x colour x gender x age). The mean of the cells "
    "that NSW would have suppressed is the estimate. QLD has no transfer-type split, so its cells are if anything slightly larger — the true NSW "
    "value may be a little lower.",
    "Fleet estimates: BEV — the value that makes growth in the NSW BEV fleet between the first and latest snapshots equal new BEV registrations "
    "(all customer types) over the same period. Other fuels — the value that makes the whole detailed snapshot add to the total in TfNSW's coarser "
    "'Age of Registered Vehicles' snapshot for the same month.",
    "Sensitivity: change the yellow cells on Inputs (e.g. set all to 1 or to 3) and every table and chart recalculates.",
]
for i, t in enumerate(notes, 12):
    wsu.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
    wsu.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
    wsu.row_dimensions[i].height = 62
widths(wsu, [16, 12, 16, 100])

# ================================================================= Charts ===
wc = wb.create_sheet("Charts")
title(wc, "Charts",
      "Native Excel charts linked to the tables — change an input and they update. Q1 = lowest-income fifth of areas (by earners), Q5 = highest.")


def group_lines(ws_src, title_, first_row, last_row, first_col, date_col=1, y_fmt=PCT, y_title=None):
    ch = LineChart()
    ch.title = title_
    for g in GROUPS:
        ch.add_data(Reference(ws_src, min_col=first_col + g - 1, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    ch.set_categories(Reference(ws_src, min_col=date_col, min_row=first_row, max_row=last_row))
    for g, s in zip(GROUPS, ch.series):
        line_series(s, GROUP_COLOURS[g - 1])
        s.tx = SeriesLabel(v=GROUP_LABEL[g])
    style_axes(ch, y_fmt, MON, y_title)
    return ch


def group_stack(ws_src, title_, first_row, last_row, first_col, date_col=1, y_fmt=NUM):
    """Stacked columns: one segment per income group, counts over time."""
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 30
    ch.title = title_
    for g in GROUPS:
        ser = Series(Reference(ws_src, min_col=first_col + g - 1, min_row=first_row, max_row=last_row), title=GROUP_LABEL[g])
        ser.graphicalProperties.solidFill = GROUP_COLOURS[g - 1]
        ser.graphicalProperties.line.solidFill = "FFFFFF"
        ch.series.append(ser)
    ch.set_categories(Reference(ws_src, min_col=date_col, min_row=first_row, max_row=last_row))
    style_axes(ch, y_fmt, MON)
    return ch


def cols(title_, ws_src, first_row, last_row, cat_col, series, y_fmt, stacked=False):
    """Column chart; series = [(col, label, colour)]."""
    ch = BarChart()
    ch.type = "col"
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
    ch.title = title_
    for c_, lab, colr in series:
        ser = Series(Reference(ws_src, min_col=c_, min_row=first_row, max_row=last_row), title=lab)
        ser.graphicalProperties.solidFill = colr
        ser.graphicalProperties.line.solidFill = "FFFFFF"
        ch.series.append(ser)
    ch.set_categories(Reference(ws_src, min_col=cat_col, min_row=first_row, max_row=last_row))
    ch.gapWidth = 40 if stacked else 60
    style_axes(ch, y_fmt, MON if stacked else None)
    return ch


def scatter(ws_src, title_, rows, xcol, ycol, colour, x_title, y_title, y_fmt=PCT, x_fmt=USD):
    ch = ScatterChart()
    ch.title = title_
    ch.style = 13
    s = Series(Reference(ws_src, min_col=ycol, min_row=rows[0], max_row=rows[1]),
               Reference(ws_src, min_col=xcol, min_row=rows[0], max_row=rows[1]), title=title_.split(" —")[0])
    s.marker.symbol = "circle"
    s.marker.size = 6
    s.marker.graphicalProperties = GraphicalProperties(solidFill=colour)
    s.marker.graphicalProperties.line.solidFill = "FFFFFF"
    s.graphicalProperties.line.noFill = True
    ch.series.append(s)
    style_axes(ch, y_fmt, x_fmt, y_title, x_title)
    ch.legend = None
    return ch


def top_bar(k):
    ttl, r1, r2, c1, c2 = top_tables[k]
    src, st = specs[k][1], specs[k][2]
    bc = BarChart()
    bc.type = "bar"
    bc.title = ttl
    ser = Series(Reference(wt, min_col=c2, min_row=r1, max_row=r2), title=ttl.split(" — ")[1])
    ser.graphicalProperties.solidFill = STATE_COLOURS["VIC" if src == "VIC" else st]
    ser.graphicalProperties.line.solidFill = "FFFFFF"
    bc.series.append(ser)
    bc.set_categories(Reference(wt, min_col=c1 + 1 if src == "VIC" else c1, min_row=r1, max_row=r2))
    bc.x_axis.scaling.orientation = "maxMin"
    bc.gapWidth = 40
    style_axes(bc, specs[k][5][-1][2])
    bc.legend = None
    return bc


GRP = (RN0, RN0 + NG - 1)   # income-group rows on Raw_Numbers
sections = [
    ("1. Raw numbers — how many BEVs", [
        cols("NSW — new private registrations per month: BEV vs other fuels", wrn, RN0, RNN, 1,
             [(3, "BEV", STATE_COLOURS["NSW"]), (4, "Other fuels", "C8C7C2")], NUM, stacked=True),
        cols("QLD — new private registrations per month: BEV vs other fuels", wrn, RN0, RNN, 1,
             [(6, "BEV", STATE_COLOURS["QLD"]), (7, "Other fuels", "C8C7C2")], NUM, stacked=True),
        cols("New BEVs registered in the recent window, by income group (count)", wrn, GRP[0], GRP[1], GC,
             [(GC + 1, "NSW", STATE_COLOURS["NSW"]), (GC + 2, "QLD", STATE_COLOURS["QLD"]),
              (GC + 3, "VIC (recent-model BEVs in fleet)", STATE_COLOURS["VIC"])], NUM),
        cols("BEVs in the registered fleet, latest, by income group (count)", wrn, GRP[0], GRP[1], GC,
             [(GC + 4, "NSW", STATE_COLOURS["NSW"]), (GC + 5, "QLD (BEVs seen since 2022)", STATE_COLOURS["QLD"]),
              (GC + 6, "VIC", STATE_COLOURS["VIC"])], NUM),
        group_stack(wk, "NSW — BEV fleet by LGA income group (count, quarter ends)", nf, nl, 2 + NG),
        group_stack(wk, "VIC — BEV fleet by postcode income group (count, quarters)", vf, vl_, 2 + NG),
    ]),
    ("2. Income — BEV share and fleet rates", [
        group_lines(wm, "NSW — BEV share of new private registrations, by LGA income group", M0, MN, SH["NSW"]),
        cols("BEV share of new vehicles, recent window, by income group", wg, grow["NSW"], grow["NSW"] + NG - 1, 2,
             [(8, "NSW (new private regos)", STATE_COLOURS["NSW"])], PCT),
        group_lines(wk, "NSW fleet — BEVs per 1,000 light vehicles, by LGA income group", nf, nl, 2 + 2 * NG, y_fmt="0", y_title="per 1,000"),
        group_lines(wk, "VIC fleet — BEVs per 1,000 vehicles, by postcode income group", vf, vl_, 2 + 4 * NG, y_fmt="0", y_title="per 1,000"),
        scatter(wl, "NSW LGAs — median income vs BEV share of new private regos (recent window)", scat["NSW"], 4, 10,
                STATE_COLOURS["NSW"], "LGA median total income ($, 2022-23)", "BEV share"),
        scatter(wp, "VIC postcodes — median taxable income vs BEVs per 1,000 vehicles", scat["VIC"], 2, 7,
                STATE_COLOURS["VIC"], "Postcode median taxable income ($, 2023-24)", "BEV per 1,000", y_fmt="0"),
    ]),
    ("3. The 2026 fuel crisis", [
        None,  # pump prices, built below
        cols("New BEVs by income group — crisis months vs same months a year earlier (count)", wrn, GRP[0], GRP[1], GC,
             [(GC + 7, "NSW year earlier", "B4B2A9"), (GC + 8, "NSW crisis", STATE_COLOURS["NSW"]),
              (GC + 9, "QLD year earlier", "D9D7CF"), (GC + 10, "QLD crisis", STATE_COLOURS["QLD"])], NUM),
        cols("NSW — BEV share by income group: crisis months vs same months a year earlier", wx, xgrow["NSW"], xgrow["NSW"] + NG - 1, 2,
             [(8, "Same months a year earlier", "B4B2A9"), (5, "Fuel-crisis months", STATE_COLOURS["NSW"])], PCT),
        cols("QLD — BEV share by income group: crisis months vs same months a year earlier", wx, xgrow["QLD"], xgrow["QLD"] + NG - 1, 2,
             [(8, "Same months a year earlier", "B4B2A9"), (5, "Fuel-crisis months", STATE_COLOURS["QLD"])], PCT),
    ]),
    ("4. Top areas (full tables on the Top10 sheet)", [top_bar(0), top_bar(2), top_bar(8), top_bar(4)]),
]
# the recent-window share bar also carries QLD and VIC
bar = sections[1][1][1]
for s_ in ("QLD", "VIC"):
    ser = Series(Reference(wg, min_col=8, min_row=grow[s_], max_row=grow[s_] + NG - 1),
                 title=f"{s_} (new private regos)" if s_ == "QLD" else "VIC (recent-model vehicles in fleet)")
    ser.graphicalProperties.solidFill = STATE_COLOURS[s_]
    ser.graphicalProperties.line.solidFill = "FFFFFF"
    bar.series.append(ser)
pr = LineChart()
pr.title = "Pump prices — the 2026 fuel crisis (retail, c/L)"
for sname in ("NSW_ULP", "NSW_Diesel", "QLD_ULP"):
    pr.add_data(Reference(wx, min_col=2 + price_cols.index(sname), min_row=X0 - 1, max_row=XN), titles_from_data=True)
pr.set_categories(Reference(wx, min_col=1, min_row=X0, max_row=XN))
for ser, colr, lab in zip(pr.series, ["1F7AE0", "9AA5B1", "F57C00"], ["NSW ULP", "NSW diesel", "QLD ULP"]):
    line_series(ser, colr)
    ser.tx = SeriesLabel(v=lab)
style_axes(pr, "0", MON, "c/L")
sections[2][1][0] = pr

ROWS_PER_CHART = 21
row = 4
for heading, chs in sections:
    c = wc.cell(row, 1, heading)
    c.font = Font(name=FONT, size=13, bold=True, color="184F95")
    row += 1
    for i, ch in enumerate(chs):
        ch.height, ch.width = 10, 17
        wc.add_chart(ch, f"{'A' if i % 2 == 0 else 'L'}{row + (i // 2) * ROWS_PER_CHART}")
    row += ((len(chs) + 1) // 2) * ROWS_PER_CHART + 1

# ================================================================ Summary ===
wsum = wb.create_sheet("Summary", 0)
title(wsum, "EV take-up by regional income — NSW, QLD, VIC",
      "Public state registration data vs ATO/ABS income by area. All numbers below are live formulas.")
head(wsum, 4, ["Measure", "NSW", "QLD", "VIC", "Note"])
gN, gQ, gV = grow["NSW"], grow["QLD"], grow["VIC"]
top, bot, all_ = NG - 1, 0, NG
fg = "Flow_Group_Summary"
fc = "Fuel_Crisis"
rn = "Raw_Numbers"
g1, g5, gA = RN0, RN0 + NG - 1, RG_ALL   # Raw_Numbers group rows
rc = lambda j: L(GC + j)                  # Raw_Numbers group-table column j
xN, xQ = xgrow["NSW"], xgrow["QLD"]
H = None   # section heading
rows = [
    ("NEW REGISTRATIONS (recent window)", H, H, H, H, None),
    ("Window", f'=TEXT({D["nsw_start"]},"mmm yyyy")&" – "&TEXT({D["nsw_last"]},"mmm yyyy")',
     f'=TEXT({D["qld_start"]},"mmm yyyy")&" – "&TEXT({D["qld_last"]},"mmm yyyy")',
     f'="Snapshot "&TEXT({D["vic_last"]},"mmm yyyy")', "VIC: no monthly regional series; recent-model vehicles in the fleet stand in", None),
    ("New private registrations (count)", f"={fg}!F{gN + all_}", f"={fg}!F{gQ + all_}", f"={fg}!F{gV + all_}",
     "VIC: recent-model vehicles in the fleet", NUM),
    ("New BEVs (count)", f"={rn}!{rc(1)}{gA}", f"={rn}!{rc(2)}{gA}", f"={rn}!{rc(3)}{gA}", None, NUM),
    ("New BEVs — lowest-income group Q1 (count)", f"={rn}!{rc(1)}{g1}", f"={rn}!{rc(2)}{g1}", f"={rn}!{rc(3)}{g1}", None, NUM),
    ("New BEVs — highest-income group Q5 (count)", f"={rn}!{rc(1)}{g5}", f"={rn}!{rc(2)}{g5}", f"={rn}!{rc(3)}{g5}", None, NUM),
    ("BEV share — all areas", f"={fg}!H{gN + all_}", f"={fg}!H{gQ + all_}", f"={fg}!H{gV + all_}", None, PCT),
    ("BEV share — Q1", f"={fg}!H{gN + bot}", f"={fg}!H{gQ + bot}", f"={fg}!H{gV + bot}", None, PCT),
    ("BEV share — Q5", f"={fg}!H{gN + top}", f"={fg}!H{gQ + top}", f"={fg}!H{gV + top}", None, PCT),
    ("BEV share Q5 ÷ Q1", f"={fg}!H{gN + NG + 1}", f"={fg}!H{gQ + NG + 1}", f"={fg}!H{gV + NG + 1}",
     "Times more likely a new car is a BEV in the richest vs poorest fifth of areas", DEC),
    ("BEV share Q5 ÷ Q1 — first 12 months of data (2023)", f"={fg}!O{gN + NG + 1}", f"={fg}!O{gQ + NG + 1}", None,
     "Lower now than in 2023 = the gap is narrowing as BEVs go mainstream", DEC),
    ("REGISTERED FLEET (latest snapshot)", H, H, H, H, None),
    ("Snapshot", f'=TEXT({D["nsw_stock"]},"mmm yyyy")', f'=TEXT({D["qld_stock"]},"mmm yyyy")', f'=TEXT({D["vic_last"]},"mmm yyyy")', None, None),
    ("BEVs in the fleet (count)", f"={rn}!{rc(4)}{gA}", f"={rn}!{rc(5)}{gA}", f"={rn}!{rc(6)}{gA}",
     "QLD: BEVs seen in any registration record since Jan 2022 (no public fleet data by region)", NUM),
    ("BEVs in the fleet — Q1 (count)", f"={rn}!{rc(4)}{g1}", f"={rn}!{rc(5)}{g1}", f"={rn}!{rc(6)}{g1}", None, NUM),
    ("BEVs in the fleet — Q5 (count)", f"={rn}!{rc(4)}{g5}", f"={rn}!{rc(5)}{g5}", f"={rn}!{rc(6)}{g5}", None, NUM),
    ("BEVs per 1,000 vehicles — all areas", f"=Stock_Group!{L(2 + 2 * NG + NG)}{nl}", "n/a", f"=Stock_Group!{L(2 + 4 * NG + NG)}{vl_}", None, "0.0"),
    ("BEVs per 1,000 vehicles — Q1", f"=Stock_Group!{L(2 + 2 * NG)}{nl}", "n/a", f"=Stock_Group!{L(2 + 4 * NG)}{vl_}", None, "0.0"),
    ("BEVs per 1,000 vehicles — Q5", f"=Stock_Group!{L(2 + 2 * NG + NG - 1)}{nl}", "n/a", f"=Stock_Group!{L(2 + 4 * NG + NG - 1)}{vl_}", None, "0.0"),
    ("2026 FUEL CRISIS (Middle East / Strait of Hormuz)", H, H, H, H, None),
    ("Crisis window (onset detected from pump prices)",
     f'=TEXT({D["cr_start"]},"mmm yyyy")&" – "&TEXT({D["cr_end"]},"mmm yyyy")',
     f'=TEXT({D["cr_start"]},"mmm yyyy")&" – "&TEXT({D["cr_end"]},"mmm yyyy")',
     f'="Quarter to "&TEXT({D["vic_last"]},"mmm yyyy")', "Compared with the same months a year earlier", None),
    ("Retail ULP: pre-crisis average → crisis peak (c/L)",
     f'=TEXT(AVERAGEIFS({fc}!B{X0}:B{XN},{fc}!A{X0}:A{XN},">="&{D["pre_start"]},{fc}!A{X0}:A{XN},"<="&{D["pre_end"]}),"0")&" → "&TEXT(_xlfn.MAXIFS({fc}!B{X0}:B{XN},{fc}!G{X0}:G{XN},1),"0")',
     f'=TEXT(AVERAGEIFS({fc}!D{X0}:D{XN},{fc}!A{X0}:A{XN},">="&{D["pre_start"]},{fc}!A{X0}:A{XN},"<="&{D["pre_end"]}),"0")&" → "&TEXT(_xlfn.MAXIFS({fc}!D{X0}:D{XN},{fc}!G{X0}:G{XN},1),"0")',
     None, None, None),
    ("New BEVs, same months a year earlier (count)", f"={fc}!G{xN + NG}", f"={fc}!G{xQ + NG}", f"={fc}!H{xgrow['VIC'] + NG}",
     "VIC: BEVs added to the fleet in the same quarter a year earlier", NUM),
    ("New BEVs, crisis months (count)", f"={fc}!D{xN + NG}", f"={fc}!D{xQ + NG}", f"={fc}!E{xgrow['VIC'] + NG}",
     "VIC: BEVs added to the fleet in the crisis quarter", NUM),
    ("New petrol/diesel/hybrid regos, change vs a year earlier", f"={fc}!N{xN + NG}", f"={fc}!N{xQ + NG}", None, None, PCT),
    ("BEV share: year earlier → crisis", f'=TEXT({fc}!H{xN + NG},"0.0%")&" → "&TEXT({fc}!E{xN + NG},"0.0%")',
     f'=TEXT({fc}!H{xQ + NG},"0.0%")&" → "&TEXT({fc}!E{xQ + NG},"0.0%")', None, None, None),
    ("BEV share multiple — Q1 (× year earlier)", f"={fc}!J{xN}", f"={fc}!J{xQ}", f"={fc}!I{xgrow['VIC']}",
     "VIC: multiple of BEVs added to the fleet", DEC),
    ("BEV share multiple — Q5 (× year earlier)", f"={fc}!J{xN + NG - 1}", f"={fc}!J{xQ + NG - 1}", f"={fc}!I{xgrow['VIC'] + NG - 1}",
     "Bigger multiple in Q1 than Q5 = the shock narrowed the gap in relative terms", DEC),
    ("BEV share change — Q1 / Q5 (percentage points)", f'=TEXT({fc}!I{xN},"0.0")&" / "&TEXT({fc}!I{xN + NG - 1},"0.0")',
     f'=TEXT({fc}!I{xQ},"0.0")&" / "&TEXT({fc}!I{xQ + NG - 1},"0.0")', None,
     "…but richer areas still added more percentage points", None),
]
for i, rr in enumerate(rows, 5):
    for j, v in enumerate(rr[:5], 1):
        c = wsum.cell(i, j, v)
        if j in (2, 3, 4) and isinstance(v, str) and v.startswith("="):
            c.font = F_LINK
            if rr[5]:
                c.number_format = rr[5]
    if rr[1] is None:
        wsum.cell(i, 1).font = F_BOLD
        for j in range(1, 6):
            wsum.cell(i, j).fill = FILL_BAND

findings_row = len(rows) + 7
wsum.cell(findings_row, 1, "How to read this").font = F_BOLD
guide = [
    "Income groups: areas are ranked by median income and cut into five groups that each hold about a fifth of the state's earners "
    "(NSW/QLD: LGAs; VIC: postcodes). Group ranks are within each state.",
    "This is an area-level (ecological) comparison: it shows where BEVs are registered, not the income of the individual buyer. "
    "Areas with many retirees or students (e.g. Sunshine Coast, Melbourne CBD) have low median taxable income but not necessarily low wealth.",
    "Novated leases (FBT-exempt for EVs since July 2022) are often registered to the employee, so they appear as private registrations "
    "and tilt BEV take-up towards areas with salaried, higher-income workers.",
    "QLD LGAs are large (Brisbane alone is about a quarter of QLD earners), so QLD groups are lumpy; NSW LGAs and VIC postcodes give a finer gradient.",
    "2026 fuel crisis: the US-Israel war with Iran from late February 2026 and the effective closure of the Strait of Hormuz sent pump "
    "prices up by around 60 c/L in March (diesel higher still in April). BEV share of new private cars jumped in both NSW and QLD from "
    "March while petrol/diesel/hybrid registrations fell. See the Fuel_Crisis sheet and section 3 of Charts.",
    "Full sources, definitions and caveats: Notes sheet. Suppression handling for NSW: Suppression sheet.",
]
for i, t in enumerate(guide, findings_row + 1):
    c = wsum.cell(i, 1, "• " + t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    wsum.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
    wsum.row_dimensions[i].height = 44
widths(wsum, [62, 16, 16, 16, 80])
wsum.freeze_panes = "A5"

# ================================================================ Sources ===
wn = wb.create_sheet("Sources")
title(wn, "Data sources — dataset pages and every file used", "All public, open data. File lists are read from raw_data/*/urls.txt.")
head(wn, 4, ["Dataset", "Publisher", "Dataset page", "Coverage used", "Licence / terms"])
datasets = [
    ("NSW Vehicle Registration Transactions (monthly)", "Transport for NSW",
     "https://opendata.transport.nsw.gov.au/data/dataset/transport-nsw-vehicle-registration-statistics",
     "Jul 2022 – latest month; new registrations by LGA, customer type, fuel", "CC BY 4.0"),
    ("NSW Vehicle Registrations Snapshot (monthly)", "Transport for NSW", "(same dataset page)",
     "Quarter-end snapshots Sep 2022 – latest; light-vehicle fleet by LGA and motive power", "CC BY 4.0"),
    ("NSW Age of Registered Vehicles Snapshot", "Transport for NSW", "(same dataset page)",
     "Dec 2025; whole-register total used to calibrate suppressed cells", "CC BY 4.0"),
    ("QLD Vehicle Registration New and Transferred Vehicle Details", "Queensland Department of Transport and Main Roads",
     "https://www.data.qld.gov.au/dataset/vehicle-registration-new-and-transfers-test", "Jan 2022 – latest; unit records by LGA", "CC BY 4.0"),
    ("VIC Whole Fleet Vehicle Registration Snapshot by Postcode", "Victorian Department of Transport and Planning",
     "https://discover.data.vic.gov.au/dataset/whole-fleet-vehicle-registration-snapshot-by-postcode", "Q2 2023 – latest quarter", "CC BY 4.0"),
    ("VIC Monthly New Vehicle Registration (checked, not used)", "Victorian Department of Transport and Planning",
     "https://discover.data.vic.gov.au/dataset/monthly-new-vehicle-registration", "No location field, so it cannot be matched to income", "CC BY 4.0"),
    ("Personal Income in Australia 2022-23, Table 1", "Australian Bureau of Statistics (ATO and other admin data)",
     "https://www.abs.gov.au/statistics/labour/earnings-and-working-conditions/personal-income-australia/latest-release",
     "Table 1.5 LGA: earners, median total income 2022-23", "CC BY 4.0"),
    ("Taxation Statistics 2023-24, Individuals Tables 6 and 8", "Australian Taxation Office",
     "https://data.gov.au/data/dataset/taxation-statistics-2023-24", "Table 8: median taxable income and individuals by postcode", "CC BY 2.5 AU"),
    ("NSW FuelCheck historical prices", "NSW Government (via au_fuel_prices project)", "https://data.nsw.gov.au/data/dataset/fuel-check",
     "Monthly average retail ULP and diesel, NSW", "CC BY 4.0"),
    ("QLD Fuel Price Reporting", "Queensland Government (via au_fuel_prices project)",
     "https://www.data.qld.gov.au/dataset/fuel-price-reporting", "Monthly average retail ULP and diesel, QLD", "CC BY 4.0"),
    ("Terminal gate prices", "Australian Institute of Petroleum (via au_fuel_prices project)",
     "https://aip.com.au/resources/historical-ulp-and-diesel-tgp-data/", "National petrol TGP, used to detect the crisis onset", "AIP terms of use"),
    ("ASGS 2023 LGA and ASGS 2021 postal area boundaries", "Australian Bureau of Statistics",
     "https://geo.abs.gov.au/arcgis/rest/services/", "Generalised polygons for the map dashboard", "CC BY 4.0"),
    ("Context: 2026 Iran war fuel crisis (ABC News)", "ABC News",
     "https://www.abc.net.au/news/2026-03-10/why-fuel-prices-are-going-up-australia/106437844", "Timing of the conflict and price shock", ""),
    ("Context: 2026 Iran war fuel crisis (Wikipedia)", "Wikipedia",
     "https://en.wikipedia.org/wiki/2026_Iran_war_fuel_crisis", "Timeline", ""),
]
for x in datasets:
    wn.append(list(x))
r = wn.max_row + 2
wn.cell(r, 1, "Individual files downloaded").font = F_BOLD
r += 1
head(wn, r, ["Dataset", "File", "URL", "", ""])
RAWD = HERE / CFG["paths"]["raw"]
url_files = [("NSW transactions", RAWD / "nsw/urls.txt"), ("NSW snapshot", RAWD / "nsw/snapshot_urls.txt"),
             ("NSW age snapshot", RAWD / "nsw/age/urls.txt"), ("QLD transactions", RAWD / "qld/urls.txt"),
             ("VIC fleet snapshot", RAWD / "vic/urls.txt"), ("Income", RAWD / "income/urls.txt")]
for lab, f in url_files:
    for line in open(f).read().split("\n"):
        u = line.split()[-1] if line.strip() else ""
        if u:
            wn.append([lab, u.split("/")[-1].split("?")[0].replace("%20", " ").replace("%2C", ","), u])
wn.append(["Fuel prices", "retail_monthly.csv, tgp_monthly.csv", f"{CFG['fuel']['retail_file']} (built by ~/models/au_fuel_prices)"])
wn.append(["Boundaries", "lga_boundaries.geojson", CFG["map"]["lga_service"]])
wn.append(["Boundaries", "vic_postcode_boundaries.geojson", CFG["map"]["poa_service"]])
for row in wn.iter_rows(min_row=5):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(c.value, str) and c.value.startswith("http"):
            c.hyperlink = c.value
            c.font = Font(name=FONT, size=10, color="1C5CAB", underline="single")
widths(wn, [44, 44, 90, 50, 22])

# ======================================================= Data adjustments ===
wa = wb.create_sheet("Data_Adjustments")
title(wa, "Summary of adjustments to the data",
      "Every filter, relabelling, imputation and proxy applied between the raw files and the tables. Counts are written by build_data.py on each run.")
head(wa, 4, ["Dataset", "Adjustment", "Detail", "Records affected"])
for x in adjustments.itertuples(index=False):
    wa.append([x.dataset, x.adjustment, x.detail, x.affected])
extra = [
    ("NSW new registrations", "Suppressed-cell value (flow)",
     "Mean '<=5' cell estimated from QLD unit records re-cut at NSW grain with a synthetic gender x age split — see Suppression sheet", ""),
    ("NSW fleet", "Suppressed-cell value (stock)",
     "BEV: calibrated so fleet growth matches new BEV registrations; other fuels: calibrated so the detailed snapshot adds to the age-snapshot total", ""),
    ("All", "Income groups within each state", "Q1–Q5 are ranked within NSW, QLD and VIC separately, so 'Q5' is relative to its own state", ""),
    ("All", "Area-level matching", "Registrations are matched to income by the owner's address area (LGA or postcode), not by the buyer's own income", ""),
    ("Fuel crisis", "Comparison windows",
     "Crisis = onset month to the latest month in both states; compared with the same calendar months a year earlier and with the same number of months just before onset", ""),
]
for x in extra:
    wa.append(list(x))
for row in wa.iter_rows(min_row=5):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
widths(wa, [24, 38, 110, 30])
wa.freeze_panes = "A5"

# ================================================================== Notes ===
wn2 = wb.create_sheet("Notes")
title(wn2, "Definitions and caveats")
head(wn2, 3, ["Type", "Item", "Detail"])
notes_rows = [
    ("Definition", "New private registration — NSW", "'Establish registration of new vehicle', customer type Private. Includes light commercials (no body-type field)."),
    ("Definition", "New private registration — QLD", "'Registration New', customer Individual, year of manufacture within one year. QLD file covers passenger body types only, "
     "so NSW and QLD levels are not directly comparable — compare gradients within each state."),
    ("Definition", "BEV", "NSW: EV / BATTERY ELECTRIC / ELECTRICITY. QLD: Electric. VIC: fuel code E. PHEV only separable in NSW."),
    ("Definition", "Fleet — NSW", "Passenger, off-road, forward-control passenger and light goods vehicles; quarter-end snapshots."),
    ("Definition", "Fleet — QLD (proxy)", "Distinct BEVs seen in any QLD transaction since Jan 2022 at their last known LGA; per 1,000 earners."),
    ("Definition", "Fleet — VIC", "Vehicle class 2 (excludes motorcycles). Recent-model = year of manufacture >= snapshot year - 1."),
    ("Definition", "Income groups", "Earner-weighted quintiles within each state. Q1 = lowest median income."),
    ("Definition", "Representation index", "Group's share of BEVs ÷ its share of all new private registrations. 1.0 = proportional."),
    ("Definition", "Top-10 tables", "Areas ranked on the metric named; only areas above the minimum-size thresholds on Inputs are ranked, so tiny areas don't top the lists on noise."),
    ("Caveat", "Area vs person", "Shows where BEVs are registered, not the buyer's income. Retiree and student areas have low median taxable income but not necessarily low wealth."),
    ("Caveat", "Novated leases", "FBT-exempt EV novated leases (since Jul 2022) are usually registered to the employee, so they count as private and favour salaried areas."),
    ("Caveat", "Income year", "LGA income 2022-23, postcode income 2023-24; registrations 2023–2026. Area income rankings move slowly."),
    ("Caveat", "QLD geography", "QLD LGAs are large (Brisbane ~ a quarter of earners), so QLD income groups are lumpy."),
    ("Caveat", "Suburbs", "No state publishes EV registrations by suburb. The finest public geography is postcode (VIC) and LGA (NSW, QLD); "
     "the 'top 10' tables use those."),
    ("Caveat", "Fuel crisis attribution", "The crisis comparison is before/after: other 2026 changes (new cheaper BEV models, NVES penalties, incentives) "
     "also act on the same months. The same-months-a-year-earlier baseline removes seasonality, not those."),
    ("Caveat", "Price sensitivity slope", "Descriptive OLS of monthly BEV share on pump price since Jan 2023; the rising BEV trend loads on it too."),
    ("Reproduce", "Scripts", "~/models/EV_income: build_data.py -> build_workbook.py -> build_dashboard.py; parameters in config.yaml"),
]
for x in notes_rows:
    wn2.append(list(x))
for row in wn2.iter_rows(min_row=4):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
widths(wn2, [11, 34, 130])

base_font(wb)
front = ["Charts", "Summary", "Raw_Numbers", "Top10", "Fuel_Crisis", "Inputs", "Flow_Group_Summary", "Flow_Group_Month",
         "Stock_Group", "LGA_Summary", "VIC_Postcode", "Suppression", "Sources", "Data_Adjustments", "Notes"]
order = [wb[n] for n in front] + [ws for ws in wb.worksheets if ws.title not in front]
wb._sheets = order
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == "Charts"
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title not in ("Summary", "Charts")
import os
out = HERE / CFG["paths"]["workbook"]
if os.environ.get("EV_PREVIEW"):
    # chart-review copy: print only the chart sheets
    out = Path(os.environ["EV_PREVIEW"])
    for ws in wb.worksheets:
        if "Chart" not in ws.title:
            ws.sheet_state = "hidden" if ws.title != "Charts" else ws.sheet_state
    for ws in wb.worksheets:
        if "Chart" in ws.title:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(out)
print("wrote", out)
